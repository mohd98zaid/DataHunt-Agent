import csv
import hashlib
import json
import os
import re
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from datahunt.config import settings
from datahunt.errors import DataHuntError, ErrorCode
from datahunt.logger import logger
from datahunt.models import ExportRecord, ExtractedRecord, ResearchRun
from datahunt.policy import escape_csv_formula
from datahunt.tools.base import ToolResult

def _discover_display_columns(records: List[ExtractedRecord], max_cols: int = 4) -> List[str]:
    """Extract up to max_cols meaningful field names across records for summary table rendering."""
    col_set: List[str] = []
    for r in records:
        for k in r.fields.keys():
            if k not in col_set and k not in ("freshness_badge", "posted_age_seconds"):
                col_set.append(k)
    return col_set[:max_cols]

def _collect_unique_sources(records: List[ExtractedRecord]) -> List[str]:
    """Collect deduplicated list of source URLs across records."""
    sources: List[str] = []
    seen = set()
    for r in records:
        if r.canonical_url and r.canonical_url not in seen:
            seen.add(r.canonical_url)
            sources.append(r.canonical_url)
    return sources

class ExportTool:
    name = "export_results"
    description = "Export verified records to CSV, JSON, or XLSX with provenance and formula injection defense."

    def __init__(self, export_dir: Optional[Path] = None):
        self.export_dir = export_dir or settings.EXPORT_DIR
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def execute(
        self,
        run_id: str,
        records: List[ExtractedRecord],
        format: str = "json",
        run_metadata: Optional[Dict[str, Any]] = None,
        columns: Optional[List[str]] = None
    ) -> ToolResult:
        format = format.lower()
        if format not in ("json", "csv", "xlsx", "md", "docx"):
            return ToolResult(
                success=False,
                error_code=ErrorCode.INVALID_REQUEST.value,
                error_message=f"Unsupported export format: {format}"
            )

        storage_key = f"export_{uuid.uuid4().hex}"
        file_name = f"{storage_key}.{format}"
        target_path = self.export_dir / file_name

        meta = run_metadata or {}
        meta["export_timestamp"] = datetime.now(timezone.utc).isoformat()
        meta["record_count"] = len(records)
        meta["run_id"] = run_id

        # Determine column list from records or provided columns
        if not columns:
            col_set = set()
            for r in records:
                col_set.update(r.fields.keys())
            # Put standard job fields first if present
            preferred_order = ["title", "company", "location", "posted_at", "application_url"]
            columns = [c for c in preferred_order if c in col_set] + sorted([c for c in col_set if c not in preferred_order])
            if not columns:
                columns = ["title", "company", "location", "posted_at", "application_url"]

        temp_path: Optional[Path] = None
        try:
            temp_fd, temp_path_str = tempfile.mkstemp(suffix=f".{format}", dir=str(self.export_dir))
            os.close(temp_fd)
            temp_path = Path(temp_path_str)

            if format == "json":
                self._write_json(temp_path, records, meta)
            elif format == "csv":
                self._write_csv(temp_path, records, columns, meta)
            elif format == "xlsx":
                self._write_xlsx(temp_path, records, columns, meta)
            elif format == "md":
                self._write_md(temp_path, records, meta)
            elif format == "docx":
                self._write_docx(temp_path, records, meta)

            # Compute SHA-256
            sha256_hash = hashlib.sha256()
            with open(temp_path, "rb") as f:
                for chunk in iter(lambda: f.read(65536), b""):
                    sha256_hash.update(chunk)
            file_hash = sha256_hash.hexdigest()

            # Atomic rename
            temp_path.replace(target_path)
            temp_path = None  # Successfully moved, clear so finally block doesn't delete

            export_record = ExportRecord(
                run_id=run_id,
                format=format,
                file_name=file_name,
                storage_key=str(target_path),
                sha256=file_hash,
                row_count=len(records)
            )

            logger.info(f"Export created successfully: {file_name} ({len(records)} rows, sha256: {file_hash[:8]}...)")
            return ToolResult(
                success=True,
                data={
                    "export_id": export_record.id,
                    "file_name": file_name,
                    "file_path": str(target_path),
                    "row_count": len(records),
                    "sha256": file_hash,
                    "format": format
                }
            )
        except Exception as e:
            logger.error(f"Failed to export results for run {run_id}: {e}")
            return ToolResult(
                success=False,
                error_code=ErrorCode.EXPORT_FAILED.value,
                error_message=f"Export generation failed: {e}"
            )
        finally:
            if temp_path and temp_path.exists():
                try:
                    temp_path.unlink(missing_ok=True)
                except Exception:
                    pass

    def _write_json(self, path: Path, records: List[ExtractedRecord], meta: Dict[str, Any]) -> None:
        data = {
            "metadata": meta,
            "records": [
                {
                    "id": r.id,
                    "fields": r.fields,
                    "canonical_url": r.canonical_url,
                    "verification_status": r.verification_status.value,
                    "confidence": r.confidence,
                    "evidence": [e.model_dump() for e in r.evidence],
                    "warnings": r.warnings,
                }
                for r in records
            ]
        }
        with open(path, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

    def _write_csv(self, path: Path, records: List[ExtractedRecord], columns: List[str], meta: Dict[str, Any]) -> None:
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            # Header
            header = columns + ["source_url", "verification_status"]
            writer.writerow(header)
            for r in records:
                row = []
                for col in columns:
                    val = r.fields.get(col)
                    row.append(escape_csv_formula(val))
                row.append(escape_csv_formula(r.canonical_url or ""))
                row.append(escape_csv_formula(r.verification_status.value))
                writer.writerow(row)

        # Write sidecar metadata JSON
        sidecar_path = path.with_suffix(".meta.json")
        with open(sidecar_path, "w", encoding="utf-8") as sf:
            json.dump(meta, sf, indent=2)

    def _write_xlsx(self, path: Path, records: List[ExtractedRecord], columns: List[str], meta: Dict[str, Any]) -> None:
        wb = openpyxl.Workbook()
        
        # Sheet 1: Results
        ws_data = wb.active
        ws_data.title = "Results"
        headers = columns + ["source_url", "verification_status"]
        ws_data.append(headers)
        
        for r in records:
            row = []
            for col in columns:
                val = r.fields.get(col)
                row.append(escape_csv_formula(val) if val is not None else "")
            row.append(escape_csv_formula(r.canonical_url or ""))
            row.append(escape_csv_formula(r.verification_status.value))
            ws_data.append(row)

        # Sheet 2: Metadata & Provenance
        ws_meta = wb.create_sheet(title="Metadata")
        ws_meta.append(["Property", "Value"])
        for k, v in meta.items():
            ws_meta.append([str(k), str(v)])

        wb.save(path)

    def _write_md(self, path: Path, records: List[ExtractedRecord], meta: Dict[str, Any]) -> None:
        lines = []
        title = meta.get("request_text") or meta.get("topic") or "Deep Technical Research Dossier"
        lines.append(f"# {title.strip()}")
        lines.append("")
        lines.append(f"> **Generated by:** {meta.get('agent_name', 'DataHunt Deep Research Agent')}")
        lines.append(f"> **Timestamp:** {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}")
        lines.append(f"> **Status:** Verified Intelligence Dossier • Sources Fetched: {meta.get('pages_fetched', 0)} • Evidence Anchors: {len(records)}")
        lines.append("")
        lines.append("---")
        lines.append("")

        # Executive Intelligence Summary
        summary = meta.get("summary", "").strip()
        if summary:
            lines.append("## Executive Intelligence Summary")
            lines.append("")
            lines.append(summary)
            lines.append("")
            lines.append("---")
            lines.append("")

        # Discovered Entities & Extracted Data
        if records:
            lines.append("## Discovered Entities & Extracted Records")
            lines.append("")
            col_set = _discover_display_columns(records, max_cols=4)
            if col_set:
                header = "| " + " | ".join(c.replace("_", " ").title() for c in col_set) + " | Status | Primary Source |"
                sep = "| " + " | ".join("---" for _ in col_set) + " | --- | --- |"
                lines.append(header)
                lines.append(sep)
                for r in records:
                    row_cells = []
                    for c in col_set:
                        val = r.fields.get(c, "")
                        val_str = str(val).replace("|", "\\|").replace("\n", " ").strip()
                        if len(val_str) > 40:
                            val_str = val_str[:37] + "..."
                        row_cells.append(val_str)
                    row_cells.append(r.verification_status.value)
                    source_link = f"[Source Link]({r.canonical_url})" if r.canonical_url else "-"
                    row_cells.append(source_link)
                    lines.append("| " + " | ".join(row_cells) + " |")
                lines.append("")
                lines.append("---")
                lines.append("")

        # Verified Source Citations
        sources = _collect_unique_sources(records)
        if sources:
            lines.append("## Primary Source Citations & References")
            lines.append("")
            for idx, s in enumerate(sources, 1):
                lines.append(f"{idx}. [{s}]({s})")
            lines.append("")

        path.write_text("\n".join(lines), encoding="utf-8")

    def _write_docx(self, path: Path, records: List[ExtractedRecord], meta: Dict[str, Any]) -> None:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        doc = docx.Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # Document Title
        title_text = meta.get("request_text") or meta.get("topic") or "Deep Technical Research Dossier"
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(title_text.strip())
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(15, 23, 42)
        title_p.paragraph_format.space_after = Pt(4)

        # Subtitle / Metadata Ribbon
        ts = meta.get("export_timestamp", "")[:10]
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(
            f"Autonomous Research Intelligence Dossier  •  Date: {ts}  •  "
            f"Sources Analyzed: {meta.get('pages_fetched', 0)}  •  Verified Entities: {meta.get('records_verified', len(records))}"
        )
        sub_run.font.size = Pt(9.5)
        sub_run.font.color.rgb = RGBColor(100, 116, 139)
        sub_run.font.italic = True
        sub_p.paragraph_format.space_after = Pt(14)

        # Executive Summary Section
        summary = (meta.get("summary") or "").strip()
        if summary:
            h1 = doc.add_heading("Executive Intelligence Summary", level=1)
            h1.paragraph_format.space_before = Pt(12)
            h1.paragraph_format.space_after = Pt(6)

            for line in summary.split("\n"):
                stripped = line.strip()
                if not stripped:
                    continue
                if stripped.startswith("### "):
                    h = doc.add_heading(stripped[4:].strip(), level=3)
                    h.paragraph_format.space_before = Pt(8)
                    h.paragraph_format.space_after = Pt(3)
                elif stripped.startswith("## "):
                    h = doc.add_heading(stripped[3:].strip(), level=2)
                    h.paragraph_format.space_before = Pt(10)
                    h.paragraph_format.space_after = Pt(4)
                elif stripped.startswith("# "):
                    h = doc.add_heading(stripped[2:].strip(), level=1)
                    h.paragraph_format.space_before = Pt(12)
                    h.paragraph_format.space_after = Pt(6)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    p = doc.add_paragraph(style='List Bullet')
                    self._add_formatted_text(p, stripped[2:])
                    p.paragraph_format.space_after = Pt(2)
                else:
                    p = doc.add_paragraph()
                    self._add_formatted_text(p, stripped)
                    p.paragraph_format.space_after = Pt(5)

        # Discovered Entities Table
        if records:
            h2 = doc.add_heading("Discovered Entities & Technical Data", level=1)
            h2.paragraph_format.space_before = Pt(16)
            h2.paragraph_format.space_after = Pt(6)

            col_set = _discover_display_columns(records, max_cols=4)
            headers = [c.replace("_", " ").title() for c in col_set] + ["Status", "Source"]

            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for idx, name in enumerate(headers):
                hdr_cells[idx].text = name
                shd = parse_xml(r'<w:shd {} w:fill="0F172A"/>'.format(nsdecls('w')))
                hdr_cells[idx]._tc.get_or_add_tcPr().append(shd)
                for run in hdr_cells[idx].paragraphs[0].runs:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9)

            for r in records:
                row_cells = table.add_row().cells
                for idx, c in enumerate(col_set):
                    val = str(r.fields.get(c, ""))
                    row_cells[idx].text = val[:45] + "..." if len(val) > 45 else val
                    if row_cells[idx].paragraphs[0].runs:
                        row_cells[idx].paragraphs[0].runs[0].font.size = Pt(8.5)
                row_cells[len(col_set)].text = r.verification_status.value
                if row_cells[len(col_set)].paragraphs[0].runs:
                    row_cells[len(col_set)].paragraphs[0].runs[0].font.size = Pt(8.5)
                row_cells[len(col_set) + 1].text = (r.canonical_url or "")[:40]
                if row_cells[len(col_set) + 1].paragraphs[0].runs:
                    row_cells[len(col_set) + 1].paragraphs[0].runs[0].font.size = Pt(8.5)

        # Primary Source Citations
        sources = _collect_unique_sources(records)
        if sources:
            h3 = doc.add_heading("Primary Source Citations & References", level=1)
            h3.paragraph_format.space_before = Pt(16)
            h3.paragraph_format.space_after = Pt(6)
            for idx, s in enumerate(sources, 1):
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(f"[{idx}] {s}")
                p.paragraph_format.space_after = Pt(2)

        doc.save(path)

    def _add_formatted_text(self, paragraph: Any, text: str) -> None:
        """Parses inline **bold** and `code` tokens and appends styled runs."""
        parts = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
        for part in parts:
            if part.startswith("**") and part.endswith("**") and len(part) >= 4:
                run = paragraph.add_run(part[2:-2])
                run.font.bold = True
            elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
                run = paragraph.add_run(part[1:-1])
                run.font.name = 'Consolas'
            else:
                paragraph.add_run(part)

